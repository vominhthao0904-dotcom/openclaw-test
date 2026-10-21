#!/usr/bin/env python3
"""Dump the structure of an existing .pptx (shapes, positions, fonts, colours) to study its layout and tone."""
import argparse, json, sys

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    sys.exit("Thiếu python-pptx: pip install python-pptx Pillow")

ap = argparse.ArgumentParser()
ap.add_argument("pptx")
ap.add_argument("--slides", help="vd: 1-5 hoặc 2,4")
ap.add_argument("--max-text", type=int, default=140)
a = ap.parse_args()
prs = Presentation(a.pptx)
pick = None
if a.slides:
    pick = set()
    for part in a.slides.split(","):
        lo, _, hi = part.partition("-")
        pick.update(range(int(lo), int(hi or lo) + 1))
out = {"slide_size_pt": [round(Emu(prs.slide_width).pt), round(Emu(prs.slide_height).pt)], "slides": []}
for i, slide in enumerate(prs.slides, 1):
    if pick and i not in pick:
        continue
    shapes = []
    for sh in slide.shapes:
        info = {"name": sh.name, "type": str(sh.shape_type), "top": round(Emu(sh.top or 0).pt, 1), "left": round(Emu(sh.left or 0).pt, 1),
                "width": round(Emu(sh.width or 0).pt, 1), "height": round(Emu(sh.height or 0).pt, 1)}
        if sh.has_text_frame and sh.text_frame.text.strip():
            info["text"] = sh.text_frame.text.strip()[: a.max_text]
            for p in sh.text_frame.paragraphs:
                if p.runs:
                    r = p.runs[0]
                    info["font"] = {"name": r.font.name, "size": r.font.size.pt if r.font.size else None, "bold": r.font.bold}
                    try:
                        if r.font.color and r.font.color.type is not None and r.font.color.rgb is not None:
                            info["font"]["color"] = str(r.font.color.rgb)
                    except AttributeError:
                        pass
                    break
        shapes.append(info)
    out["slides"].append({"index": i, "layout": slide.slide_layout.name, "shapes": shapes})
print(json.dumps(out, ensure_ascii=False, indent=2))
