#!/usr/bin/env python3
import argparse
import io
import sys
import json
import math
import os
import re
import tempfile
from pathlib import Path
from urllib.parse import urlparse, urlencode

import ipaddress
import socket
import urllib.request
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

SLIDE_W = 13.333
SLIDE_H = 7.5

WHITE = RGBColor(255, 255, 255)
INK = RGBColor(18, 24, 38)
NAVY = RGBColor(16, 44, 84)
NAVY_2 = RGBColor(31, 64, 122)
RED = RGBColor(220, 38, 38)
GOLD = RGBColor(245, 158, 11)
SKY = RGBColor(14, 165, 233)
TEAL = RGBColor(20, 184, 166)
PURPLE = RGBColor(124, 58, 237)
PINK = RGBColor(236, 72, 153)
SLATE = RGBColor(100, 116, 139)
BORDER = RGBColor(207, 216, 228)
SOFT = RGBColor(244, 247, 251)
SOFT_2 = RGBColor(232, 240, 252)

ACCENTS = [RED, GOLD, SKY, TEAL, PURPLE, PINK]
USED_IMAGE_KEYS = set()
WARNINGS = []
FOOTER_TEXT = ''
MAX_IMAGE_BYTES = 15 * 1024 * 1024


def load_slides(path: str):
    with open(path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    if isinstance(data, dict) and 'slides' in data:
        return data['slides']
    if isinstance(data, list):
        return data
    raise ValueError('JSON input phải là list slides hoặc object chứa key "slides"')


def ensure_dir(path: Path):
    path.mkdir(parents=True, exist_ok=True)
    return path


def clean_filename(text: str):
    return re.sub(r'[^a-zA-Z0-9_-]+', '-', text).strip('-').lower() or 'image'


def _host_is_public(host: str) -> bool:
    try:
        infos = socket.getaddrinfo(host, None)
    except OSError:
        return False
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            return False
    return True


def download_file(url: str, out_dir: Path, stem: str):
    """HTTPS only, public hosts only, size-capped, must decode as an image."""
    u = urlparse(url)
    if u.scheme != 'https' or not u.hostname:
        raise ValueError('chỉ chấp nhận ảnh https')
    if not _host_is_public(u.hostname):
        raise ValueError('host không công khai (chặn truy cập mạng nội bộ)')
    req = urllib.request.Request(url, headers={'User-Agent': 'pptx-deck-builder/1.0'})
    with urllib.request.urlopen(req, timeout=30) as r:
        data = r.read(MAX_IMAGE_BYTES + 1)
    if len(data) > MAX_IMAGE_BYTES:
        raise ValueError('ảnh quá lớn (>15MB)')
    out = out_dir / f'{clean_filename(stem)}.img'
    out.write_bytes(data)
    with Image.open(out) as im:
        fmt = (im.format or 'JPEG').lower()
    final = out.with_suffix('.png' if fmt == 'png' else '.webp' if fmt == 'webp' else '.jpg')
    out.rename(final)
    return final


def resolve_image_source(slide_data):
    src = slide_data.get('image_path') or slide_data.get('image_url')
    return str(src) if src else None


def download_image_if_needed(source: str, out_dir: Path, stem: str):
    if not source:
        return None
    if source.startswith('http://') or source.startswith('https://'):
        try:
            return str(download_file(source, out_dir, stem))
        except Exception as e:
            WARNINGS.append(f"Không tải được ảnh '{stem}': {e}")
            return None
    if os.path.exists(source):
        return source
    WARNINGS.append(f"Không thấy file ảnh: {source}")
    return None


def add_textbox(slide, x, y, w, h, text='', font_size=20, color=INK, bold=False,
               align=PP_ALIGN.LEFT, name=None, fill=None, line=None, margin=0.1, valign=MSO_ANCHOR.TOP,
               allow_empty=False):
    if (text is None or str(text).strip() == '') and not allow_empty:
        return None
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin * 0.8)
    tf.margin_bottom = Inches(margin * 0.5)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = 'Arial'
    if fill:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    else:
        box.fill.background()
    if line:
        box.line.color.rgb = line
    else:
        box.line.fill.background()
    return box


def add_bullets(slide, x, y, w, h, items, font_size=18, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.space_after = Pt(16)
        p.line_spacing = 1.18
        p.alignment = PP_ALIGN.LEFT
        for run in p.runs:
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
            run.font.name = 'Arial'
    box.line.fill.background()
    box.fill.background()
    return box


def add_rect(slide, x, y, w, h, fill, line=BORDER, radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()
    return shape


def crop_to_fill(image_path: str, target_ratio: float):
    img = Image.open(image_path).convert('RGB')
    w, h = img.size
    current = w / h
    if current > target_ratio:
        new_w = int(h * target_ratio)
        left = (w - new_w) // 2
        box = (left, 0, left + new_w, h)
    else:
        new_h = int(w / target_ratio)
        top = (h - new_h) // 2
        box = (0, top, w, top + new_h)
    return img.crop(box)


def add_cover_image(slide, image_path: str, x, y, w, h, temp_dir: Path):
    target_ratio = w / h
    cropped = crop_to_fill(image_path, target_ratio)
    out = temp_dir / f'crop-{Path(image_path).stem}-{abs(hash((x, y, w, h)))}.jpg'
    cropped.save(out, quality=92)
    slide.shapes.add_picture(str(out), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    return str(out)


def add_footer(slide):
    if not FOOTER_TEXT:
        return None
    bar = add_rect(slide, 0, 7.0, SLIDE_W, 0.5, NAVY, NAVY, radius=False)
    add_textbox(slide, 8.2, 7.03, 4.7, 0.32, FOOTER_TEXT, 12, WHITE, True, PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)
    return bar


def add_title(slide, title: str, eyebrow: str = None):
    if eyebrow:
        add_textbox(slide, 0.58, 0.28, 4.8, 0.26, eyebrow.upper(), 10, RED, True)
    add_textbox(slide, 0.55, 0.48, 12.0, 0.74, title, 24, NAVY, True)
    add_rect(slide, 0.55, 1.16, 1.25, 0.05, RED, None, radius=False)


def add_kpi_row(slide, metrics):
    if not metrics:
        return
    box_w = 2.45
    gap = 0.22
    start_x = 0.55
    y = 1.45
    for idx, metric in enumerate(metrics[:4]):
        x = start_x + idx * (box_w + gap)
        fill = [SOFT, SOFT_2, RGBColor(240, 253, 250), RGBColor(250, 245, 255)][idx % 4]
        add_rect(slide, x, y, box_w, 0.95, fill, BORDER)
        add_textbox(slide, x + 0.12, y + 0.08, box_w - 0.24, 0.28, str(metric.get('label', 'KPI')), 11, SLATE, True)
        add_textbox(slide, x + 0.12, y + 0.30, box_w - 0.24, 0.45, str(metric.get('value', '0')), 22, NAVY, True)
        if metric.get('note'):
            add_textbox(slide, x + 0.12, y + 0.65, box_w - 0.24, 0.18, str(metric['note']), 9, RED)


def _claim_image(local, title):
    key = str(Path(local).resolve())
    if key in USED_IMAGE_KEYS:
        WARNINGS.append(f"Ảnh bị dùng lặp ở slide '{title}': {Path(local).name}")
    USED_IMAGE_KEYS.add(key)


def add_photo_block(slide, slide_data, side: str, temp_dir: Path):
    image_source = resolve_image_source(slide_data)
    title = slide_data.get('title', 'Untitled')
    local = download_image_if_needed(image_source, temp_dir, title) if image_source else None
    if not local:
        WARNINGS.append(f"Slide '{title}' không có ảnh: đã chuyển sang bố cục chữ toàn chiều ngang")
        add_bullets(slide, 0.85, 2.0, 11.5, 4.2, slide_data.get('content', []), 20, INK)
        return
    if side == 'right':
        img_x, txt_x = 7.05, 0.72
    else:
        img_x, txt_x = 0.72, 6.88
    img_y, img_w, img_h = 2.2, 5.48, 4.05
    add_bullets(slide, txt_x, 2.2, 5.25, 3.55, slide_data.get('content', []), 16, INK)
    _claim_image(local, title)
    add_cover_image(slide, local, img_x, img_y, img_w, img_h, temp_dir)


def add_cards(slide, slide_data):
    items = slide_data.get('cards') or [{'title': c, 'body': ''} for c in slide_data.get('content', [])]
    items = items[:4]
    count = max(1, len(items))
    total_gap = 0.26 * (count - 1)
    card_w = (11.9 - total_gap) / count
    x = 0.70
    colors = [SOFT, SOFT_2, RGBColor(239, 246, 255), RGBColor(250, 245, 255)]
    for idx, card in enumerate(items):
        add_rect(slide, x, 2.02, card_w, 3.60, colors[idx % len(colors)], BORDER)
        add_rect(slide, x + 0.18, 2.24, 0.58, 0.08, ACCENTS[idx % len(ACCENTS)], None, radius=False)
        add_textbox(slide, x + 0.18, 2.42, card_w - 0.36, 0.58, str(card.get('title', 'Card')), 17, NAVY, True)
        body = card.get('body') or card.get('text') or ''
        if isinstance(body, list):
            body = ' • '.join(body)
        add_textbox(slide, x + 0.18, 3.00, card_w - 0.36, 1.82, body, 10, INK)
        x += card_w + 0.26


def add_center_statement(slide, slide_data, temp_dir: Path):
    image_source = resolve_image_source(slide_data)
    title = slide_data.get('title', 'hero')
    local = download_image_if_needed(image_source, temp_dir, title) if image_source else None
    statement = slide_data.get('statement') or slide_data.get('title')
    subtitle = slide_data.get('subtitle') or ' • '.join(slide_data.get('content', [])[:3])

    if local:
        add_rect(slide, 0.9, 1.75, 11.55, 2.2, SOFT, BORDER)
        _claim_image(local, title)
        add_cover_image(slide, local, 0.9, 1.75, 11.55, 2.2, temp_dir)
        overlay = add_rect(slide, 0.9, 1.75, 11.55, 2.2, NAVY, None, radius=False)
        overlay.fill.transparency = 0.45
        add_rect(slide, 1.15, 4.2, 11.0, 1.85, SOFT, BORDER)
        add_textbox(slide, 1.55, 4.45, 10.2, 0.78, statement, 24, NAVY, True,
                    PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if subtitle:
            add_textbox(slide, 2.0, 5.25, 9.3, 0.45, subtitle, 12, INK, False, PP_ALIGN.CENTER)
        add_rect(slide, 4.75, 5.72, 3.85, 0.18, RED, None)
        return

    # Không có ảnh: đừng để một khối màu rỗng chiếm nửa slide — cho câu chốt đứng giữa,
    # chữ to hơn. Đây là dáng mặc định khi chưa có khoá ảnh, nên nó phải tự nó đẹp.
    WARNINGS.append(f"Slide '{title}' (statement) không có ảnh: dùng bố cục chữ giữa trang")
    add_rect(slide, 0.9, 2.05, 11.55, 3.5, SOFT, BORDER)
    add_textbox(slide, 1.5, 2.6, 10.35, 1.5, statement, 34, NAVY, True,
                PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_textbox(slide, 2.1, 4.25, 9.15, 0.6, subtitle, 14, INK, False, PP_ALIGN.CENTER)
    add_rect(slide, 5.75, 5.05, 1.85, 0.14, RED, None)


def add_process_flow(slide, slide_data):
    steps = slide_data.get('steps') or slide_data.get('content', [])[:5]
    if not steps:
        steps = ['Bước 1', 'Bước 2', 'Bước 3']
    count = len(steps)
    step_w = 2.2 if count <= 4 else 1.9
    gap = 0.22
    total = count * step_w + (count - 1) * gap
    start_x = max(0.6, (12.5 - total) / 2)
    y = 2.55
    for idx, step in enumerate(steps):
        x = start_x + idx * (step_w + gap)
        add_rect(slide, x, y, step_w, 1.55, [SOFT, SOFT_2, RGBColor(240, 253, 250), RGBColor(250, 245, 255), RGBColor(254, 242, 242)][idx % 5], BORDER)
        add_textbox(slide, x + 0.16, y + 0.15, 0.5, 0.3, str(idx + 1), 13, RED, True)
        add_textbox(slide, x + 0.16, y + 0.45, step_w - 0.32, 0.72, str(step), 16, NAVY, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if idx < count - 1:
            connector = add_rect(slide, x + step_w, y + 0.68, gap, 0.12, SKY, None, radius=False)
            connector.line.fill.background()


def add_chart_slide(slide, slide_data):
    chart = slide_data.get('chart') or {}
    categories = chart.get('categories') or []
    values = chart.get('values') or []
    if not categories or len(categories) != len(values):
        raise SystemExit(f"Slide '{slide_data.get('title', 'Untitled')}' (chart) thiếu dữ liệu thật: cần chart.categories và chart.values cùng độ dài. "
                         "Script không tự bịa số liệu: hãy xin người dùng cung cấp.")
    series_name = chart.get('series_name') or 'Giá trị'
    chart_kind = (chart.get('type') or 'column').strip().lower()

    chart_type_map = {
        'column': XL_CHART_TYPE.COLUMN_CLUSTERED,
        'bar': XL_CHART_TYPE.BAR_CLUSTERED,
        'line': XL_CHART_TYPE.LINE_MARKERS,
        'area': XL_CHART_TYPE.AREA,
        'pie': XL_CHART_TYPE.PIE,
        'doughnut': XL_CHART_TYPE.DOUGHNUT,
    }
    ppt_chart_type = chart_type_map.get(chart_kind, XL_CHART_TYPE.COLUMN_CLUSTERED)

    add_rect(slide, 0.68, 1.92, 7.95, 4.35, WHITE, BORDER)
    add_rect(slide, 0.68, 1.92, 7.95, 0.16, NAVY, None, radius=False)
    data = CategoryChartData()
    data.categories = categories
    data.add_series(series_name, values)
    chart_shape = slide.shapes.add_chart(
        ppt_chart_type,
        Inches(0.98), Inches(2.28), Inches(7.25), Inches(3.35), data
    )
    chart_obj = chart_shape.chart
    chart_obj.has_legend = chart_kind in {'pie', 'doughnut', 'line', 'area'}
    if chart_obj.has_legend:
        chart_obj.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart_obj.legend.font.size = Pt(9)
        chart_obj.legend.include_in_layout = False
    chart_obj.chart_title.has_text_frame = False

    fill_colors = [NAVY, SKY, TEAL, GOLD, RED, PURPLE]

    if chart_kind not in {'pie', 'doughnut'}:
        chart_obj.value_axis.has_major_gridlines = True
        chart_obj.value_axis.major_gridlines.format.line.color.rgb = RGBColor(232, 236, 243)
        chart_obj.category_axis.tick_labels.font.size = Pt(9)
        chart_obj.value_axis.tick_labels.font.size = Pt(9)

    if hasattr(chart_obj, 'plots') and chart_obj.plots:
        plot = chart_obj.plots[0]
        if hasattr(plot, 'gap_width') and chart_kind in {'column', 'bar'}:
            plot.gap_width = 58

    if chart_kind in {'column', 'bar', 'pie', 'doughnut'}:
        points = chart_obj.series[0].points
        for idx, point in enumerate(points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = fill_colors[idx % len(fill_colors)]
            point.format.line.color.rgb = fill_colors[idx % len(fill_colors)]
    else:
        series = chart_obj.series[0]
        series.format.line.color.rgb = NAVY
        series.format.line.width = Pt(2.5)
        series.smooth = True if chart_kind == 'line' else False
        if chart_kind == 'area':
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = SKY
            series.format.fill.transparency = 0.20

    insight_text = slide_data.get('insight')
    if insight_text:
        add_rect(slide, 8.92, 1.92, 3.68, 1.92, RGBColor(247, 250, 252), BORDER)
        add_textbox(slide, 9.16, 2.16, 3.10, 0.24, slide_data.get('insight_title') or 'Nhận định', 12, RED, True)
        add_textbox(slide, 9.16, 2.48, 3.06, 0.96, insight_text, 10, INK)

    stats = slide_data.get('stats') or []
    y = 4.10 if insight_text else 1.92
    stat_fills = [SOFT_2, RGBColor(240, 253, 250), RGBColor(255, 247, 237)]
    for idx, stat in enumerate(stats[:3]):
        add_rect(slide, 8.92, y, 3.68, 0.68, stat_fills[idx % len(stat_fills)], BORDER)
        add_textbox(slide, 9.12, y + 0.08, 1.70, 0.18, str(stat.get('label', '')), 8, SLATE, True)
        add_textbox(slide, 10.76, y + 0.02, 1.50, 0.26, str(stat.get('value', '')), 16, NAVY, True, PP_ALIGN.RIGHT)
        y += 0.80


def add_quote_proof(slide, slide_data, temp_dir: Path):
    add_rect(slide, 0.85, 1.95, 8.0, 3.85, SOFT, BORDER)
    quote = slide_data.get('quote') or slide_data.get('statement') or slide_data.get('title')
    add_textbox(slide, 1.2, 2.3, 7.25, 2.0, f'“{quote}”', 24, NAVY, True)
    if slide_data.get('quote_author'):
        add_textbox(slide, 1.2, 4.6, 6.0, 0.35, slide_data['quote_author'], 11, RED, True)
    image_source = slide_data.get('image_path') or slide_data.get('image_url')
    if image_source:
        local = download_image_if_needed(str(image_source), temp_dir, slide_data.get('title', 'proof'))
        if local:
            add_rect(slide, 9.15, 2.0, 3.05, 1.55, SOFT, BORDER)
            add_cover_image(slide, local, 9.15, 2.0, 3.05, 1.55, temp_dir)
    proof = slide_data.get('proofs') or slide_data.get('content', [])[:3]
    y = 3.72 if image_source else 2.1
    for idx, item in enumerate(proof[:2 if image_source else 3]):
        add_rect(slide, 9.15, y, 3.05, 0.9, [SOFT_2, RGBColor(240, 253, 250), RGBColor(250, 245, 255)][idx % 3], BORDER)
        add_textbox(slide, 9.4, y + 0.16, 2.55, 0.36, item, 12, INK, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        y += 1.02


def render_slide(prs, slide_data, temp_dir: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    layout = (slide_data.get('layout') or 'split-right').strip().lower()
    add_title(slide, slide_data.get('title', 'Untitled'), slide_data.get('eyebrow'))
    if layout not in {'split-right', 'split-left', 'split'}:
        add_kpi_row(slide, slide_data.get('metrics'))

    if layout in {'split-right', 'split'}:
        add_photo_block(slide, slide_data, 'right', temp_dir)
    elif layout == 'split-left':
        add_photo_block(slide, slide_data, 'left', temp_dir)
    elif layout in {'cards', 'grid'}:
        add_cards(slide, slide_data)
    elif layout in {'center', 'statement'}:
        add_center_statement(slide, slide_data, temp_dir)
    elif layout in {'process', 'timeline', 'flow'}:
        add_process_flow(slide, slide_data)
    elif layout in {'chart', 'data'}:
        add_chart_slide(slide, slide_data)
    elif layout in {'quote', 'proof'}:
        add_quote_proof(slide, slide_data, temp_dir)
    else:
        add_bullets(slide, 0.85, 2.0, 11.5, 4.2, slide_data.get('content', []), 20, INK)

    add_footer(slide)
    return slide


def create_pitch_deck(slides_data, output_path='output.pptx'):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    temp_dir = ensure_dir(Path(tempfile.mkdtemp(prefix='pptx-deck-')))
    for slide_data in slides_data:
        render_slide(prs, slide_data, temp_dir)

    prs.save(output_path)
    return output_path


def audit_empty_text(path):
    prs = Presentation(path)
    bad = []
    for i, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            if sh.has_text_frame and sh.shape_type == 17 and not sh.text_frame.text.strip():  # 17 = text box
                bad.append(i)
                break
    return bad


def main():
    parser = argparse.ArgumentParser(description='Build PowerPoint deck from JSON spec')
    parser.add_argument('--input', required=True, help='JSON file chứa slides')
    parser.add_argument('--output', required=True, help='Đường dẫn file .pptx output')
    parser.add_argument('--footer', default='', help='Chữ chân trang (mặc định: không có chân trang)')
    args = parser.parse_args()

    global FOOTER_TEXT
    FOOTER_TEXT = args.footer.strip()
    slides = load_slides(args.input)
    out = create_pitch_deck(slides, args.output)
    empties = audit_empty_text(out)
    for w in WARNINGS:
        print(f'[CẢNH BÁO] {w}', file=sys.stderr)
    for e in empties:
        print(f'[CẢNH BÁO] Slide {e}: có ô chữ rỗng', file=sys.stderr)
    print(json.dumps({'ok': True, 'file': str(out), 'slides': len(slides), 'warnings': len(WARNINGS) + len(empties)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
